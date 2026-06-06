"""PPTX slide deck generation for agent eval artifacts."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any


def generate_slides_pptx(output_path: Path, title: str, slides: list[dict[str, Any]]) -> dict[str, Any]:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("python-pptx is required for slide generation") from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()

    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = "Generated inside SkillVault TEE"

    bullet_layout = prs.slide_layouts[1]
    for item in slides[:12]:
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = str(item.get("title", "Slide"))
        body = slide.shapes.placeholders[1].text_frame
        body.clear()
        bullets = item.get("bullets") or item.get("points") or []
        if isinstance(bullets, str):
            bullets = [bullets]
        for idx, bullet in enumerate(bullets[:6]):
            p = body.paragraphs[0] if idx == 0 else body.add_paragraph()
            p.text = str(bullet)
            p.level = 0
            p.font.size = Pt(18)

    prs.save(str(output_path))
    size = output_path.stat().st_size
    return {"path": str(output_path), "size": size, "slide_count": len(slides)}


def run_generate_slides_tool(args: dict[str, Any], workspace: Path) -> str:
    title = str(args.get("title", "Research Talk"))
    slides = args.get("slides") or []
    if not isinstance(slides, list) or not slides:
        slides = [
            {"title": "Problem", "bullets": ["Motivation from buyer papers"]},
            {"title": "Approach", "bullets": ["Mechanism and evaluation"]},
            {"title": "Takeaways", "bullets": ["Summary and open questions"]},
        ]
    output = workspace / "output" / "slides.pptx"
    result = generate_slides_pptx(output, title, slides)
    return json.dumps(result)
